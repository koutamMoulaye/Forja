"""Deterministic generation of the sample datasets used by the test-suite.

Run ``uv run python -m tests.fixtures.generate`` to rebuild them.
"""

from __future__ import annotations

import csv
import io
import json
import random
from datetime import date, datetime, timedelta
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook

FIXTURES_DIR = Path(__file__).parent

FR_CSV = "clients_fr_sale.csv"
EXCEL = "ventes.xlsx"
PARQUET = "produits.parquet"
JSON_ARRAY = "evenements.json"
JSON_LINES = "evenements.jsonl"
EMPTY = "vide.csv"
HEADER_ONLY = "entete_seule.csv"
CORRUPTED_PARQUET = "corrompu.parquet"
CORRUPTED_EXCEL = "corrompu.xlsx"
BINARY_CSV = "binaire.csv"
DUPLICATED_COLUMNS = "colonnes_dupliquees.csv"
UTF8_COMMA_CSV = "orders_utf8.csv"

PEOPLE: list[tuple[str, str]] = [
    ("Jean", "Dupont"),
    ("Marie", "Martin"),
    ("Pierre", "Bernard"),
    ("Sophie", "Petit"),
    ("Nicolas", "Durand"),
    ("Camille", "Leroy"),
    ("Julien", "Moreau"),
    ("Léa", "Simon"),
    ("Thomas", "Laurent"),
    ("Chloé", "Lefèvre"),
    ("Antoine", "Michel"),
    ("Manon", "Garcia"),
    ("Louis", "David"),
    ("Emma", "Bertrand"),
    ("Hugo", "Roux"),
    ("Inès", "Vincent"),
    ("Arthur", "Fournier"),
    ("Jade", "Morel"),
    ("Gabriel", "Girard"),
    ("Louise", "André"),
    ("Raphaël", "Mercier"),
    ("Alice", "Blanc"),
    ("Paul", "Guérin"),
    ("Claire", "Boyer"),
    ("Mathieu", "Garnier"),
    ("Julie", "Chevalier"),
    ("Olivier", "François"),
    ("Isabelle", "Legrand"),
    ("Vincent", "Gauthier"),
    ("Nathalie", "Perrin"),
]

FR_ROW_COUNT = len(PEOPLE)
FR_DUPLICATED_INDEXES = (2, 6)
FR_MISSING_EMAIL_INDEXES = (4, 17)
FR_INVALID_AMOUNT_INDEX = 11


def _ascii(value: str) -> str:
    import unicodedata

    normalized = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in normalized if not unicodedata.combining(ch))


def fr_iban(bank: int, branch: int, account: int, key: int) -> str:
    bban = f"{bank:05d}{branch:05d}{account:011d}{key:02d}"
    numeric = "".join(str(int(ch, 36)) for ch in bban + "FR00")
    check = 98 - int(numeric) % 97
    iban = f"FR{check:02d}{bban}"
    return " ".join(iban[i : i + 4] for i in range(0, len(iban), 4))


def _fr_amount(rng: random.Random, index: int) -> str:
    if index == FR_INVALID_AMOUNT_INDEX:
        return "N/A"
    value = round(rng.uniform(0.5, 25_000), 2)
    integer, cents = f"{value:.2f}".split(".")
    if len(integer) > 3 and index % 2 == 0:
        integer = f"{integer[:-3]} {integer[-3:]}"
    return f"{integer},{cents}"


def _fr_phone(rng: random.Random, index: int) -> str:
    digits = f"{rng.randint(10_000_000, 99_999_999):08d}"
    prefix = rng.choice(["6", "7", "1"])
    match index % 4:
        case 0:
            return f"0{prefix} {digits[0:2]} {digits[2:4]} {digits[4:6]} {digits[6:8]}"
        case 1:
            return f"+33 {prefix} {digits[0:2]} {digits[2:4]} {digits[4:6]} {digits[6:8]}"
        case 2:
            return f"0{prefix}{digits}"
        case _:
            return f"0{prefix}.{digits[0:2]}.{digits[2:4]}.{digits[4:6]}.{digits[6:8]}"


def fr_csv_rows() -> list[list[str]]:
    rng = random.Random(42)
    postal_codes = ["75001", "69002", "13008", "01000", "06000", "33000", "59000", "02100"]
    comments = ["Client fidèle", "Très satisfait", "à relancer", "Appel; rappel demain", ""]
    rows: list[list[str]] = []
    for index, (first, last) in enumerate(PEOPLE):
        signup = date(2022, 1, 1) + timedelta(days=rng.randint(0, 900))
        email = (
            ""
            if index in FR_MISSING_EMAIL_INDEXES
            else f"{_ascii(first).lower()}.{_ascii(last).lower()}@example.fr"
        )
        rows.append(
            [
                str(index + 1),
                last,
                first,
                email,
                _fr_phone(rng, index),
                rng.choice(postal_codes),
                signup.strftime("%d/%m/%Y"),
                _fr_amount(rng, index),
                fr_iban(30006, 10000 + index, 12345678900 + index * 7, 10 + index),
                rng.choice(["oui", "non"]),
                comments[index % len(comments)],
            ]
        )
    rows.extend(list(rows[i]) for i in FR_DUPLICATED_INDEXES)
    return rows


FR_HEADER = [
    "id",
    "nom",
    "prénom",
    "email",
    "téléphone",
    "code_postal",
    "date_inscription",
    "montant_total",
    "iban",
    "actif",
    "commentaire",
]


def write_fr_csv(path: Path) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(FR_HEADER)
    writer.writerows(fr_csv_rows())
    path.write_bytes(buffer.getvalue().encode("cp1252"))


def write_excel(path: Path) -> None:
    rng = random.Random(7)
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Commandes"
    orders.append(["order_id", "line_no", "product_id", "quantity", "unit_price", "ordered_at"])
    start = datetime(2024, 3, 1, 9, 0)
    for order_id in range(1, 21):
        for line_no in range(1, rng.randint(2, 4)):
            orders.append(
                [
                    order_id,
                    line_no,
                    rng.randint(1, 10),
                    rng.randint(1, 5),
                    round(rng.uniform(2, 80), 2),
                    start + timedelta(hours=order_id * 5),
                ]
            )
    stores = workbook.create_sheet("Magasins")
    stores.append(["store_id", "city", "opened_on"])
    for store_id, city in enumerate(["Paris", "Lyon", "Marseille", "Lille"], start=1):
        stores.append([store_id, city, date(2015 + store_id, 5, 1)])
    workbook.create_sheet("Vide")
    workbook.save(path)


def write_parquet(path: Path) -> None:
    labels = [f"Produit {i:02d}" for i in range(1, 11)]
    table = pa.table(
        {
            "product_id": pa.array(range(1, 11), type=pa.int64()),
            "label": labels,
            "category": ["outillage", "jardin"] * 5,
            "price": [round(3.5 * i, 2) for i in range(1, 11)],
            "created_at": [datetime(2023, 1, i) for i in range(1, 11)],
        }
    )
    pq.write_table(table, path)


def _events() -> list[dict[str, object]]:
    return [
        {
            "event_id": i,
            "type": "click" if i % 2 else "view",
            "occurred_at": f"2024-05-{i:02d}T10:15:00",
            "user": {"id": 100 + i % 3, "email": f"user{i % 3}@example.com"},
            "tags": ["web", "fr"] if i % 2 else [],
        }
        for i in range(1, 9)
    ]


def write_json_array(path: Path) -> None:
    path.write_text(json.dumps(_events(), indent=2), encoding="utf-8")


def write_json_lines(path: Path) -> None:
    path.write_text("\n".join(json.dumps(e) for e in _events()) + "\n", encoding="utf-8")


def write_invalid_files(directory: Path) -> None:
    (directory / EMPTY).write_bytes(b"")
    (directory / HEADER_ONLY).write_text("id;nom;email\r\n", encoding="utf-8")
    (directory / CORRUPTED_PARQUET).write_bytes(b"PAR1" + bytes(range(256)) * 4 + b"PAR1")
    (directory / CORRUPTED_EXCEL).write_bytes(b"PK\x03\x04" + b"not really a workbook" * 20)
    rng = random.Random(3)
    (directory / BINARY_CSV).write_bytes(bytes(rng.randint(0, 255) for _ in range(4096)))
    (directory / DUPLICATED_COLUMNS).write_text(
        "id,Name,name,,value,id\n1,a,b,c,10,1\n2,d,e,f,20,2\n", encoding="utf-8"
    )
    (directory / UTF8_COMMA_CSV).write_text(
        'order_id,customer,amount,ordered_on\n1,"Smith, John",1234.50,2024-01-15\n'
        "2,Zoë Kraus,99.90,2024-02-01\n3,Ana Lima,10.00,2024-02-03\n",
        encoding="utf-8",
    )


def generate_all(directory: Path = FIXTURES_DIR) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    write_fr_csv(directory / FR_CSV)
    write_excel(directory / EXCEL)
    write_parquet(directory / PARQUET)
    write_json_array(directory / JSON_ARRAY)
    write_json_lines(directory / JSON_LINES)
    write_invalid_files(directory)


ALL_FILES = (
    FR_CSV,
    EXCEL,
    PARQUET,
    JSON_ARRAY,
    JSON_LINES,
    EMPTY,
    HEADER_ONLY,
    CORRUPTED_PARQUET,
    CORRUPTED_EXCEL,
    BINARY_CSV,
    DUPLICATED_COLUMNS,
    UTF8_COMMA_CSV,
)


if __name__ == "__main__":
    generate_all()
