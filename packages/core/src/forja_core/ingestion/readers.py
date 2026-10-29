"""Readers turning uploaded files into raw Arrow tables (one per sheet for Excel)."""

from __future__ import annotations

import csv
import io
import json
import sys
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq

from forja_core.config import DuplicateColumnPolicy
from forja_core.errors import (
    CorruptedFileError,
    DuplicateColumnsError,
    EmptyFileError,
    ForjaError,
)
from forja_core.ingestion.detection import decode_text, detect_delimiter
from forja_core.models import TECHNICAL_COLUMNS, SourceFormat

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))


@dataclass
class RawTable:
    table: pa.Table
    sheet_name: str | None = None
    encoding: str | None = None
    delimiter: str | None = None
    warnings: list[str] = field(default_factory=list)


def normalize_headers(
    raw_names: Sequence[object], policy: DuplicateColumnPolicy
) -> tuple[list[str], list[str]]:
    """Make column names non-empty, unique (case-insensitively) and non-reserved."""
    names: list[str] = []
    warnings: list[str] = []
    taken: set[str] = set()
    duplicates: list[str] = []
    reserved = {name.lower() for name in TECHNICAL_COLUMNS}

    for index, raw in enumerate(raw_names, start=1):
        name = "" if raw is None else str(raw).strip()
        if not name:
            name = f"column_{index}"
            warnings.append(f"column {index} has no header, renamed to '{name}'")
        if name.lower() in reserved:
            renamed = f"{name}_src"
            warnings.append(f"column '{name}' uses a reserved name, renamed to '{renamed}'")
            name = renamed
        if name.lower() in taken:
            duplicates.append(name)
            base = name
            suffix = 2
            while f"{base}_{suffix}".lower() in taken:
                suffix += 1
            name = f"{base}_{suffix}"
            warnings.append(f"duplicated column '{base}' renamed to '{name}'")
        taken.add(name.lower())
        names.append(name)

    if duplicates and policy is DuplicateColumnPolicy.ERROR:
        raise DuplicateColumnsError(
            "source contains duplicated column names",
            details={"columns": sorted(set(duplicates))},
        )
    return names, warnings


def _cell_to_text(value: Any) -> str:
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, dict | list):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def to_arrow_array(values: list[Any]) -> pa.Array:
    """Build a typed array when values share a scalar type, otherwise fall back to text."""
    non_null = [v for v in values if v is not None]
    if not non_null:
        return pa.array(values, type=pa.string())
    kinds = {type(v) for v in non_null}
    arrow_type: pa.DataType | None = None
    if kinds == {bool}:
        arrow_type = pa.bool_()
    elif kinds == {int}:
        arrow_type = pa.int64()
    elif kinds <= {int, float}:
        arrow_type = pa.float64()
    elif kinds == {datetime}:
        arrow_type = pa.timestamp("us")
    elif kinds == {date}:
        arrow_type = pa.date32()
    elif kinds == {str}:
        arrow_type = pa.string()
    if arrow_type is not None:
        try:
            if arrow_type == pa.float64():
                return pa.array([None if v is None else float(v) for v in values], type=arrow_type)
            if arrow_type == pa.timestamp("us"):
                values = [v.replace(tzinfo=None) if isinstance(v, datetime) else v for v in values]
            return pa.array(values, type=arrow_type)
        except (pa.ArrowInvalid, pa.ArrowTypeError, OverflowError):
            pass
    return pa.array([None if v is None else _cell_to_text(v) for v in values], type=pa.string())


def _build_table(names: list[str], columns: list[list[Any]]) -> pa.Table:
    return pa.Table.from_arrays([to_arrow_array(col) for col in columns], names=names)


def read_csv(path: Path, policy: DuplicateColumnPolicy) -> RawTable:
    raw = path.read_bytes()
    if not raw.strip():
        raise EmptyFileError("file is empty")
    text, encoding = decode_text(raw)
    text = text.lstrip("﻿")
    if not text.strip():
        raise EmptyFileError("file is empty")

    delimiter = detect_delimiter(text)
    try:
        rows = [
            row
            for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
            if any(cell.strip() for cell in row)
        ]
    except csv.Error as exc:
        raise CorruptedFileError(f"malformed CSV content: {exc}") from exc
    if not rows:
        raise EmptyFileError("file is empty")
    header, body = rows[0], rows[1:]
    if not body:
        raise EmptyFileError("file contains a header but no data rows")

    warnings: list[str] = []
    width = max(len(header), *(len(row) for row in body))
    if width > len(header):
        warnings.append(
            f"some rows have more fields than the header ({width} > {len(header)}), "
            "extra columns were added"
        )
        header = header + [""] * (width - len(header))
    short_rows = sum(1 for row in body if len(row) < width)
    if short_rows:
        warnings.append(f"{short_rows} row(s) have missing trailing fields, padded with nulls")

    names, header_warnings = normalize_headers(header, policy)
    warnings = header_warnings + warnings
    columns: list[list[Any]] = [[] for _ in names]
    for row in body:
        for index in range(width):
            value = row[index] if index < len(row) else None
            columns[index].append(value if value is not None and value.strip() else None)

    table = pa.Table.from_arrays([pa.array(col, type=pa.string()) for col in columns], names=names)
    return RawTable(table=table, encoding=encoding, delimiter=delimiter, warnings=warnings)


def read_excel(path: Path, policy: DuplicateColumnPolicy) -> list[RawTable]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    if path.stat().st_size == 0:
        raise EmptyFileError("file is empty")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise CorruptedFileError(f"cannot open Excel workbook: {exc}") from exc

    tables: list[RawTable] = []
    skipped: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                list(row)
                for row in sheet.iter_rows(values_only=True)
                if any(cell is not None and str(cell).strip() for cell in row)
            ]
            if len(rows) < 2:
                skipped.append(sheet.title)
                continue
            header, body = rows[0], rows[1:]
            width = max(len(header), *(len(row) for row in body))
            header = header + [None] * (width - len(header))
            # Drop trailing columns without header nor data.
            while (
                width > 1
                and header[width - 1] is None
                and all(len(row) < width or row[width - 1] is None for row in body)
            ):
                width -= 1
            names, warnings = normalize_headers(header[:width], policy)
            columns: list[list[Any]] = [[] for _ in names]
            for row in body:
                for index in range(width):
                    value = row[index] if index < len(row) else None
                    if isinstance(value, str) and not value.strip():
                        value = None
                    columns[index].append(value)
            tables.append(
                RawTable(
                    table=_build_table(names, columns), sheet_name=sheet.title, warnings=warnings
                )
            )
    finally:
        workbook.close()

    if not tables:
        raise EmptyFileError("workbook contains no sheet with data")
    if skipped:
        message = f"empty sheet(s) skipped: {', '.join(skipped)}"
        for raw_table in tables:
            raw_table.warnings.append(message)
    return tables


def read_parquet(path: Path, policy: DuplicateColumnPolicy) -> RawTable:
    if path.stat().st_size == 0:
        raise EmptyFileError("file is empty")
    try:
        table = pq.read_table(path)
    except (pa.ArrowInvalid, pa.ArrowException, OSError) as exc:
        raise CorruptedFileError(f"cannot read Parquet file: {exc}") from exc
    if table.num_rows == 0:
        raise EmptyFileError("Parquet file contains no rows")
    names, warnings = normalize_headers(table.column_names, policy)
    return RawTable(table=table.rename_columns(names), warnings=warnings)


def _flatten(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    flat: dict[str, Any] = {}
    for key, value in record.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict) and value:
            flat.update(_flatten(value, prefix=f"{name}."))
        elif isinstance(value, dict | list):
            flat[name] = json.dumps(value, ensure_ascii=False, default=str)
        else:
            flat[name] = value
    return flat


def _parse_json_records(text: str) -> tuple[list[Any], list[str]]:
    warnings: list[str] = []
    try:
        document = json.loads(text)
    except json.JSONDecodeError:
        records: list[Any] = []
        for line_number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise CorruptedFileError(
                    f"invalid JSON at line {line_number}: {exc.msg}",
                    details={"line": line_number},
                ) from exc
        return records, warnings

    if isinstance(document, list):
        return document, warnings
    if isinstance(document, dict):
        list_values = [(k, v) for k, v in document.items() if isinstance(v, list)]
        if len(list_values) == 1 and all(isinstance(item, dict) for item in list_values[0][1]):
            key, records_list = list_values[0]
            warnings.append(f"records extracted from top-level key '{key}'")
            return records_list, warnings
        return [document], warnings
    return [{"value": document}], warnings


def read_json(path: Path, policy: DuplicateColumnPolicy) -> RawTable:
    raw = path.read_bytes()
    if not raw.strip():
        raise EmptyFileError("file is empty")
    text, encoding = decode_text(raw)
    text = text.lstrip("﻿").strip()
    if not text:
        raise EmptyFileError("file is empty")

    records, warnings = _parse_json_records(text)
    if not records:
        raise EmptyFileError("JSON document contains no records")

    flattened = [_flatten(r) if isinstance(r, dict) else {"value": r} for r in records]
    keys: list[str] = []
    seen: set[str] = set()
    for record in flattened:
        for key in record:
            if key not in seen:
                seen.add(key)
                keys.append(key)
    names, header_warnings = normalize_headers(keys, policy)
    columns = [[record.get(key) for record in flattened] for key in keys]
    return RawTable(
        table=_build_table(names, columns),
        encoding=encoding,
        warnings=header_warnings + warnings,
    )


def read_source(path: Path, fmt: SourceFormat, policy: DuplicateColumnPolicy) -> list[RawTable]:
    try:
        match fmt:
            case SourceFormat.CSV:
                return [read_csv(path, policy)]
            case SourceFormat.EXCEL:
                return read_excel(path, policy)
            case SourceFormat.PARQUET:
                return [read_parquet(path, policy)]
            case SourceFormat.JSON:
                return [read_json(path, policy)]
            case _:
                raise CorruptedFileError(f"{fmt.value} sources are not files")
    except ForjaError:
        raise
    except Exception as exc:
        raise CorruptedFileError(f"cannot read {fmt.value} file: {exc}") from exc
